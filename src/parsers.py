"""
parsers.py
----------
Extracts raw text from resumes and job descriptions in PDF, DOCX, or
plain-text format. This is the "input" stage of the agent pipeline.
"""

from pathlib import Path

import pdfplumber
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(file_path: str) -> str:
    """Read a resume/JD file and return its plain-text content.

    Supports .pdf, .docx, and .txt/.md files. Raises ValueError for
    anything else so the caller can skip/report it instead of crashing
    the whole batch run.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    # .txt / .md
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_pdf(path: Path) -> str:
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            chunks.append(text)
    return "\n".join(chunks)


def _extract_docx(path: Path) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]

    # Tables often hold skills/experience in resume templates - grab those too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def collect_resume_files(folder: str) -> list[str]:
    """Return a sorted list of supported resume file paths in a folder."""
    folder_path = Path(folder)
    if not folder_path.is_dir():
        raise NotADirectoryError(f"Resume folder not found: {folder}")

    files = [
        str(p)
        for p in sorted(folder_path.iterdir())
        if p.suffix.lower() in SUPPORTED_EXTENSIONS and p.is_file()
    ]
    return files
